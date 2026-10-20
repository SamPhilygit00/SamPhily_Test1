#!/usr/bin/env python3
"""Extract the TPS/TVQ paid on expenses (purchases, service fees...) from the
documents dropped in the repo's depenses/ folder.

Where to put the documents:
  depenses/AAAA-TN/...   e.g. depenses/2026-T4/facture_shopify_octobre.pdf
                         every document in a quarter folder counts for that
                         quarter, whatever the dates written on it.
  depenses/...           a document directly in depenses/ (not in a quarter
                         folder) counts each expense in the quarter of its date.

Supported documents:
  .xlsx / .xlsm / .csv   a table with one row per expense. The header row must
                         have a TPS (or GST/TVH/HST) column and a TVQ (or QST)
                         column; Date, Fournisseur, Description, N° facture,
                         Montant HT / Sous-total and Total columns are used when
                         present. Read directly, no AI involved.
  .pdf                   invoices / receipts (one or several per file).
  .jpg .jpeg .png .webp .gif .heic
                         photos or scans of receipts.

PDFs and images are read by Claude through the Anthropic API (environment
variable ANTHROPIC_API_KEY). The result of each document is cached in
data/depenses/cache.json by content hash, so a document is only sent once.
Without ANTHROPIC_API_KEY, a PDF that contains text is read with a simple
pattern search (flagged "à vérifier") and images are reported as not read.
A table without recognisable TPS/TVQ columns is also sent to Claude when the
key is set.

Taxes included: when a document only shows a total "taxes incluses" (or TTC,
tax included) without the TPS/TVQ amounts, they are worked out from that
total at the Québec rates: HT = total / 1.14975, TPS = 5 % of HT, TVQ = the
rest (so HT + TPS + TVQ = total exactly). The expense is flagged "calculée".
Putting "taxes-incluses" (or "TTC") in a file name forces this calculation
for that file, e.g. for a receipt that doesn't say so; in a table, it
applies to the rows that have a total but no TPS/TVQ.

Combined tax (no TPS/TVQ split printed, but a tax amount IS shown, e.g.
"Taxes de vente", "Taxes estimées"): split from that amount using the same
5 % / 9,975 % rate ratio (split_combined_tax). A handwritten "?" next to
"TPS" and/or "TVQ" on the document (e.g. "TPS - ?", "TVQ - ?") is the
convention used to flag this on documents dropped in depenses/ — it tells
Claude to compute the tax instead of leaving it at 0, overriding the usual
"ignore handwritten annotations" rule for that one marker.

Batching: multiple invoices in a single PDF (e.g. a stack of receipts
scanned together) are all extracted as separate expenses from one API call.
A very large deposit (many invoices in one PDF) should still be split into
several smaller PDFs of under ~30 MB each (GitHub's upload/attachment
limit) — see depenses/README.md for how.

Usage (standalone, prints the quarter's expenses and totals):
  python scripts/extract_expenses.py 2026-T4

Optional environment variables:
  ANTHROPIC_API_KEY  enables reading PDFs and photos
  ANTHROPIC_MODEL    default "claude-sonnet-5-5"
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import sys
import time
import unicodedata
from datetime import date, datetime
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parent.parent
EXPENSES_DIR = ROOT / "depenses"
CACHE_PATH = ROOT / "data" / "depenses" / "cache.json"

QUARTER_DIR_RE = re.compile(r"^(\d{4})-[Tt]([1-4])$")
TABLE_EXT = {".xlsx", ".xlsm", ".csv"}
PDF_EXT = {".pdf"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".heic", ".heif"}
IGNORED_NAMES = {".gitkeep", "readme.md", ".ds_store", "thumbs.db"}

API_URL = "https://api.anthropic.com/v1/messages"
DEFAULT_MODEL = "claude-sonnet-5-5"
MAX_IMAGE_BYTES = 1_500_000  # photos above this are resized (API limit: 5 MB per image)
CACHE_VERSION = 5  # 5: max_tokens raised (multi-invoice PDFs were truncating the tool call)

TPS_RATE = 0.05
TVQ_RATE = 0.09975
TAX_INCLUDED_NAME_RE = re.compile(r"taxes?[ _.-]*inclus|(?<![a-z])ttc(?![a-z])|tax[ _.-]*incl", re.IGNORECASE)
TAX_INCLUDED_TEXT_RE = re.compile(r"taxes?\s+inclus|\bTTC\b|tax(es)?\s+incl", re.IGNORECASE)


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def normalize(text) -> str:
    """Lower-case, strip accents and punctuation: 'N° Facture' -> 'n facture'."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


CATEGORY_DEFAULT = "Matériel de cils"

# Checked in order against the normalized supplier name; first match wins.
CATEGORY_RULES: list[tuple[re.Pattern, str]] = [
    (re.compile(r"\bamazon\b"), "Matériel de cils"),
    (re.compile(r"\bdhl\b"), "Frais de livraison"),
    (re.compile(r"\bscores\b"), "Frais de maintien divers"),
    (re.compile(r"\bmaxi\b"), "Frais de maintien divers"),
    (re.compile(r"\bdollarama\b"), "Frais de maintien divers"),
    (re.compile(r"\bh\s?m\b"), "Frais de maintien divers"),
    (re.compile(r"\bdollar\s+max\b"), "Frais de maintien divers"),
    (re.compile(r"\bmichaels\b"), "Frais de maintien divers"),
    (re.compile(r"\bla perle virtuelle\b"), "Soutien technique"),
    (re.compile(r"\bzahka\b"), "Service d'accompagnement en ligne de marque"),
    (re.compile(r"\btelus\b"), "Frais Téléphonie et Prêt"),
    (re.compile(r"\bmm\b"), "Matériel de cils"),
]


def categorize_fournisseur(fournisseur: str) -> str:
    text = normalize(fournisseur)
    for pattern, category in CATEGORY_RULES:
        if pattern.search(text):
            return category
    return CATEGORY_DEFAULT


def parse_amount(value) -> float | None:
    """'1 234,56 $' -> 1234.56, '(12.00)' -> -12.0, 12 -> 12.0, '' -> None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(" ", "").replace(" ", "")
    text = text.replace("$", "").replace("CAD", "").replace("€", "")
    if not text or text in {"-", "—"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        amount = float(text)
    except ValueError:
        return None
    return -amount if negative else amount


def parse_date(value) -> str | None:
    """Return an ISO date (AAAA-MM-JJ) or None. Day-first for 01/02/2026."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        for candidate in (text, text[:10]):
            try:
                return datetime.strptime(candidate, fmt).date().isoformat()
            except ValueError:
                continue
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}-{match.group(3)}"
    return None


def quarter_of_iso(iso: str | None) -> str | None:
    if not iso:
        return None
    year, month = int(iso[:4]), int(iso[5:7])
    return f"{year}-T{(month - 1) // 3 + 1}"


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def relative(path: Path) -> str:
    try:
        return path.relative_to(EXPENSES_DIR).as_posix()
    except ValueError:
        return path.name


def r2(value) -> float:
    return round(float(value or 0), 2)


def split_tax_included(total: float) -> tuple[float, float, float]:
    """Total with TPS and TVQ included -> (HT, TPS, TVQ), summing exactly to total."""
    ht = round(total / (1 + TPS_RATE + TVQ_RATE), 2)
    tps = round(ht * TPS_RATE, 2)
    tvq = round(total - ht - tps, 2)
    return ht, tps, tvq


def apply_tax_included(record: dict) -> dict:
    """Fill TPS/TVQ from the total when the expense shows no tax amounts."""
    if record["tps"] or record["tvq"] or not record["total"] or record["devise"] != "CAD":
        return record
    ht, tps, tvq = split_tax_included(record["total"])
    record.update(montant_ht=ht, tps=tps, tvq=tvq)
    record["remarques"] = record["remarques"] + [
        f"TPS/TVQ calculées à partir du total taxes incluses ({record['total']:.2f} $ = "
        f"{ht:.2f} HT + {tps:.2f} TPS 5 % + {tvq:.2f} TVQ 9,975 %)"]
    record["methode"] += ", taxes calculées"
    return record


def split_combined_tax(taxe: float) -> tuple[float, float]:
    """A single printed tax amount (TPS+TVQ not itemized separately) -> (TPS, TVQ),
    split by the fixed 5 % / 9,975 % rate ratio so TPS + TVQ sums exactly to taxe."""
    tps = round(taxe * TPS_RATE / (TPS_RATE + TVQ_RATE), 2)
    tvq = round(taxe - tps, 2)
    return tps, tvq


def apply_combined_tax(record: dict, taxe_combinee) -> dict:
    """Fill TPS/TVQ from a single combined tax amount printed on the document
    (e.g. "Taxes de vente", "Taxes estimées") when TPS and TVQ are not itemized."""
    taxe = parse_amount(taxe_combinee)
    if not taxe or record["tps"] or record["tvq"] or record["devise"] != "CAD":
        return record
    tps, tvq = split_combined_tax(taxe)
    record.update(tps=tps, tvq=tvq)
    record["remarques"] = record["remarques"] + [
        f"TPS/TVQ réparties à partir de la taxe combinée imprimée ({taxe:.2f} $ = "
        f"{tps:.2f} TPS 5 % + {tvq:.2f} TVQ 9,975 %)"]
    record["methode"] += ", taxe combinée répartie"
    return record


def new_record(fichier: str, methode: str, **fields) -> dict:
    tps = r2(fields.get("tps"))
    tvh = r2(fields.get("tvh"))
    remarques = list(fields.get("remarques") or [])
    if tvh:
        remarques.append(f"TVH de {tvh:.2f} $ comptée avec la TPS")
    return {
        "fichier": fichier,
        "date": fields.get("date"),
        "fournisseur": (fields.get("fournisseur") or "").strip(),
        "facture": str(fields.get("facture") or "").strip(),
        "description": (fields.get("description") or "").strip(),
        "montant_ht": r2(fields.get("montant_ht")) if fields.get("montant_ht") is not None else None,
        "tps": r2(tps + tvh),
        "tvq": r2(fields.get("tvq")),
        "total": r2(fields.get("total")) if fields.get("total") is not None else None,
        "devise": (fields.get("devise") or "CAD").upper(),
        "methode": methode,
        "remarques": remarques,
    }


# --------------------------------------------------------------------------
# Tables (xlsx / csv)
# --------------------------------------------------------------------------

# Checked in this order: the first matching role wins for a header cell.
COLUMN_ROLES = [
    ("tvq", [r"\btvq\b", r"\bqst\b"]),
    ("tvh", [r"\btvh\b", r"\bhst\b"]),
    ("tps", [r"\btps\b", r"\bgst\b"]),
    ("date", [r"\bdate\b"]),
    ("facture", [r"\bfacture\b", r"\binvoice\b", r"\bn\b", r"\bnumero\b", r"\breference\b", r"\bref\b", r"\bno\b"]),
    ("fournisseur", [r"\bfournisseur\b", r"\bvendeur\b", r"\bmarchand\b", r"\bcommerce\b", r"\bsupplier\b",
                     r"\bvendor\b", r"\bmagasin\b", r"\bnom\b"]),
    ("montant_ht", [r"\bsous total\b", r"\bsubtotal\b", r"\bht\b", r"\bavant tax", r"\bhors tax", r"\bbrut\b",
                    r"\bmontant\b"]),
    ("total", [r"\btotal\b", r"\bttc\b"]),
    ("description", [r"\bdescription\b", r"\blibelle\b", r"\bdetail\b", r"\bobjet\b", r"\barticle\b",
                     r"\bcategorie\b", r"\bnature\b"]),
]


def header_roles(row: list) -> dict[str, int]:
    roles: dict[str, int] = {}
    for index, cell in enumerate(row):
        text = normalize(cell)
        if not text:
            continue
        for role, patterns in COLUMN_ROLES:
            if role in roles:
                continue
            if any(re.search(p, text) for p in patterns):
                roles[role] = index
                break
    return roles


def table_records(rows: list[list], fichier: str, sheet: str = "") -> list[dict] | None:
    """Parse a table; None when no header row with TPS and TVQ columns is found."""
    header_index = None
    roles: dict[str, int] = {}
    for i, row in enumerate(rows[:25]):
        candidate = header_roles(row)
        if "tvq" in candidate and ("tps" in candidate or "tvh" in candidate):
            header_index, roles = i, candidate
            break
    if header_index is None:
        return None

    label = f"tableau{f' ({sheet})' if sheet else ''}"
    tax_included = bool(TAX_INCLUDED_NAME_RE.search(fichier))
    records = []
    for row in rows[header_index + 1:]:
        def get(role):
            idx = roles.get(role)
            return row[idx] if idx is not None and idx < len(row) else None

        first_text = next((normalize(c) for c in row if normalize(c)), "")
        if not first_text:
            continue
        if first_text.startswith(("total", "sous total", "grand total")):
            continue
        tps, tvq, tvh = parse_amount(get("tps")), parse_amount(get("tvq")), parse_amount(get("tvh"))
        ht, total = parse_amount(get("montant_ht")), parse_amount(get("total"))
        if all(v is None for v in (tps, tvq, tvh, ht, total)):
            continue
        record = new_record(
            fichier, label,
            date=parse_date(get("date")),
            fournisseur=str(get("fournisseur") or ""),
            facture=get("facture"),
            description=str(get("description") or ""),
            montant_ht=ht, tps=tps, tvh=tvh, tvq=tvq, total=total,
        )
        if tax_included and tps is None and tvq is None and tvh is None:
            record = apply_tax_included(record)
        records.append(record)
    return records


def read_table_rows(path: Path) -> list[tuple[str, list[list]]]:
    """Return [(sheet name, rows)] for an xlsx or csv file."""
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [("", [row for row in csv.reader(io.StringIO(text), dialect)])]

    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
    wb.close()
    return sheets


def table_as_text(sheets: list[tuple[str, list[list]]], limit: int = 60000) -> str:
    out = []
    for name, rows in sheets:
        if name:
            out.append(f"### Feuille : {name}")
        for row in rows:
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)[:limit]


# --------------------------------------------------------------------------
# Claude (PDF, images, unrecognised tables)
# --------------------------------------------------------------------------

TOOL = {
    "name": "enregistrer_depenses",
    "description": "Enregistre chaque facture ou reçu trouvé dans le document, avec les taxes payées.",
    "input_schema": {
        "type": "object",
        "properties": {
            "depenses": {
                "type": "array",
                "description": "Une entrée par facture ou reçu distinct (pas par article).",
                "items": {
                    "type": "object",
                    "properties": {
                        "date": {"type": ["string", "null"], "description": "Date de la facture, AAAA-MM-JJ"},
                        "fournisseur": {"type": "string"},
                        "facture": {"type": ["string", "null"], "description": "Numéro de facture ou de reçu"},
                        "description": {"type": "string", "description": "Ce qui a été acheté, en quelques mots"},
                        "montant_ht": {"type": ["number", "null"], "description": "Montant avant taxes"},
                        "tps": {"type": "number", "description": "TPS / GST payée (0 si absente)"},
                        "tvh": {"type": "number", "description": "TVH / HST payée (0 si absente)"},
                        "tvq": {"type": "number", "description": "TVQ / QST payée (0 si absente)"},
                        "total": {"type": ["number", "null"], "description": "Total payé, taxes incluses"},
                        "devise": {"type": "string", "description": "Code de devise, ex. CAD, USD"},
                        "taxes_incluses": {"type": "boolean",
                                           "description": "true si le document indique que le total inclut "
                                                          "les taxes (« taxes incluses », « TTC », « tax "
                                                          "included », ou un reçu/total final sans AUCUNE "
                                                          "mention de taxe) SANS imprimer le moindre montant "
                                                          "de taxe"},
                        "taxe_combinee": {"type": ["number", "null"],
                                          "description": "Montant de taxe imprimé quand TPS et TVQ ne sont "
                                                         "PAS indiqués séparément mais qu'un montant de taxe "
                                                         "EST visible et distinct du total (ex. « Taxes de "
                                                         "vente : 77,46 $ », « Taxes estimées : 83,11 $ », "
                                                         "« Taxe : 4,49 $ »). Utilise le montant de taxe du "
                                                         "TOTAL FINAL réellement payé (après rabais), pas "
                                                         "celui d'une ligne d'article avant rabais. Laisse "
                                                         "tps=0 et tvq=0 dans ce cas ; null si TPS/TVQ sont "
                                                         "indiqués séparément ou si taxes_incluses est true"},
                        "remarque": {"type": ["string", "null"],
                                     "description": "Doute ou particularité (illisible, remboursement, devise étrangère...)"},
                    },
                    "required": ["date", "fournisseur", "facture", "description", "montant_ht",
                                 "tps", "tvh", "tvq", "total", "devise", "taxes_incluses",
                                 "taxe_combinee", "remarque"],
                    "additionalProperties": False,
                },
            },
            "document_lisible": {"type": "boolean"},
        },
        "required": ["depenses", "document_lisible"],
        "additionalProperties": False,
    },
}

PROMPT = (
    "Ce document contient une ou plusieurs dépenses d'une petite entreprise du Québec "
    "(factures, reçus, frais de service). Pour CHAQUE facture ou reçu distinct, relève les "
    "montants de taxes réellement inscrits sur le document : TPS (GST), TVH (HST) et TVQ (QST). "
    "Ne calcule jamais une taxe qui n'est pas imprimée : si une taxe n'apparaît pas, mets 0. "
    "TROIS cas possibles selon ce qui est imprimé, à distinguer soigneusement : "
    "(1) TPS et TVQ sont indiquées séparément -> mets ces montants dans tps/tvq, taxes_incluses=false, "
    "taxe_combinee=null. "
    "(2) un seul montant de taxe est imprimé, SANS distinction TPS/TVQ, mais ce montant de taxe est "
    "visible et distinct du total (ex. « Taxes de vente : 77,46 $ », « Taxes estimées : 83,11 $ », "
    "« Taxe : 4,49 $ ») -> mets ce montant dans taxe_combinee, tps=0, tvq=0, taxes_incluses=false. "
    "(3) aucun montant de taxe n'apparaît du tout, seulement un total final payé (c'est le cas normal "
    "d'un reçu de service comme Uber, livraison, etc., où la taxe est incluse sans être détaillée) -> "
    "mets ce total dans total, tps=0, tvq=0, taxes_incluses=true, taxe_combinee=null. "
    "Chaque dépense doit correspondre à exactement UN de ces trois cas. Utilise toujours les montants du "
    "total final réellement payé (après rabais/réduction), jamais ceux d'une ligne d'article avant rabais, "
    "même si cette ligne affiche une TPS/TVQ détaillée. "
    "Ignore les annotations écrites à la main (crochets, flèches, notes, catégories) : seuls les montants "
    "imprimés par le marchand comptent pour les CHIFFRES. SEULE EXCEPTION : un « ? » écrit à la main à côté "
    "de « TPS » et/ou « TVQ » (ex. « TPS - ? », « TVQ - ? ») est un repère intentionnel de l'entreprise "
    "signalant que cette taxe n'est pas détaillée sur le document et doit être calculée : traite alors ce "
    "cas comme (2) ou (3) ci-dessus selon ce qui est imprimé (taxe combinée visible, ou total seulement), "
    "jamais comme une raison de laisser la dépense de côté. "
    "Ne confonds pas les numéros d'inscription (ex. « TPS : 123456789 RT0001 ») avec des montants. "
    "Un remboursement ou une note de crédit a des montants négatifs. "
    "Si le document est illisible ou n'est pas une dépense, renvoie une liste vide et document_lisible=false. "
    "Réponds uniquement en appelant l'outil enregistrer_depenses."
)


def _post(payload: dict, headers: dict) -> requests.Response:
    for attempt in range(4):
        resp = requests.post(API_URL, json=payload, headers=headers, timeout=180)
        if resp.status_code in (429, 500, 502, 503, 529) and attempt < 3:
            time.sleep(float(resp.headers.get("retry-after", 5 * (attempt + 1))))
            continue
        return resp
    return resp


def _json_from_text(text: str) -> dict | None:
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        data = json.loads(text[start:end + 1])
    except json.JSONDecodeError:
        return None
    return data if isinstance(data, dict) and "depenses" in data else None


def call_claude(content_block: dict, api_key: str) -> dict:
    """Ask Claude to fill the enregistrer_depenses tool.

    Recent models (e.g. Claude Sonnet 5.5) reject forced tool use
    (tool_choice "tool"/"any"), so the call uses tool_choice "auto" with a
    strict tool definition, which guarantees schema-valid inputs."""
    model = os.environ.get("ANTHROPIC_MODEL") or DEFAULT_MODEL
    tool = dict(TOOL, strict=True)
    payload = {
        "model": model,
        "max_tokens": 16000,
        "tools": [tool],
        "tool_choice": {"type": "auto"},
        "messages": [{"role": "user", "content": [content_block, {"type": "text", "text": PROMPT}]}],
    }
    headers = {"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"}
    resp = _post(payload, headers)
    if resp.status_code == 400 and "strict" in resp.text:
        payload["tools"] = [TOOL]  # model without strict tool support
        resp = _post(payload, headers)
    if resp.status_code >= 400:
        raise RuntimeError(f"API Anthropic {resp.status_code} : {resp.text[:300]}")

    body = resp.json()
    stop_reason = body.get("stop_reason")
    content = body.get("content", [])
    for block in content:
        if block.get("type") == "tool_use" and block.get("name") == TOOL["name"]:
            result = block.get("input") or {}
            if "depenses" not in result:
                # max_tokens cut the response before the tool call was complete.
                raise RuntimeError(f"Appel d'outil incomplet (stop_reason={stop_reason}) : "
                                   "augmenter max_tokens ou réduire la taille du document")
            result["_modele"] = model
            return result
    text = "".join(b.get("text", "") for b in content if b.get("type") == "text")
    result = _json_from_text(text)
    if result is not None:
        result["_modele"] = model
        return result
    raise RuntimeError("Réponse de l'API sans résultat structuré : " + text[:200])


def image_block(path: Path) -> dict:
    from PIL import Image, ImageOps

    suffix = path.suffix.lower()
    if suffix in (".heic", ".heif"):
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError as exc:
            raise RuntimeError("format HEIC non pris en charge (pillow-heif manquant)") from exc

    data = path.read_bytes()
    media = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
             ".webp": "image/webp", ".gif": "image/gif"}.get(suffix)
    if media is None or len(data) > MAX_IMAGE_BYTES:
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
        img.thumbnail((2400, 2400))
        quality = 90
        while True:
            buffer = io.BytesIO()
            img.save(buffer, format="JPEG", quality=quality)
            if buffer.tell() <= MAX_IMAGE_BYTES or quality <= 50:
                break
            quality -= 10
        data, media = buffer.getvalue(), "image/jpeg"
    return {"type": "image", "source": {"type": "base64", "media_type": media,
                                        "data": base64.b64encode(data).decode()}}


def records_from_claude(result: dict, fichier: str, methode: str) -> list[dict]:
    forced = bool(TAX_INCLUDED_NAME_RE.search(fichier))
    records = []
    for item in result.get("depenses") or []:
        remarques = [item["remarque"]] if item.get("remarque") else []
        devise = (item.get("devise") or "CAD").upper()
        if devise != "CAD":
            remarques.append(f"Devise {devise} : montants non convertis, vérifier")
        records.append(new_record(
            fichier, methode,
            date=parse_date(item.get("date")),
            fournisseur=item.get("fournisseur"),
            facture=item.get("facture"),
            description=item.get("description"),
            montant_ht=item.get("montant_ht"),
            tps=item.get("tps"), tvh=item.get("tvh"), tvq=item.get("tvq"),
            total=item.get("total"), devise=devise, remarques=remarques,
        ))
        if item.get("taxes_incluses") or forced:
            records[-1] = apply_tax_included(records[-1])
        elif item.get("taxe_combinee"):
            records[-1] = apply_combined_tax(records[-1], item["taxe_combinee"])
    return records


# --------------------------------------------------------------------------
# PDF fallback without API key
# --------------------------------------------------------------------------

AMOUNT = r"(-?\d{1,3}(?:[  ]?\d{3})*[.,]\d{2})"


def pdf_text_records(path: Path, fichier: str) -> list[dict]:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    if not text.strip():
        return []

    def last_amount(labels: str) -> float:
        found = 0.0
        for line in text.splitlines():
            if re.search(labels, line, re.IGNORECASE) and not re.search(r"\bRT\s?\d{4}|\bTQ\s?\d{4}|\d{9,}", line):
                amounts = re.findall(AMOUNT, line)
                if amounts:
                    found = parse_amount(amounts[-1]) or 0.0
        return found

    tps = last_amount(r"\b(TPS|GST)\b")
    tvh = last_amount(r"\b(TVH|HST)\b")
    tvq = last_amount(r"\b(TVQ|QST)\b")
    total = last_amount(r"\bTOTAL\b")
    dates = re.findall(r"\b(\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4})\b", text)
    included = TAX_INCLUDED_TEXT_RE.search(text) or TAX_INCLUDED_NAME_RE.search(fichier)
    if not (tps or tvh or tvq) and not (included and total):
        return []
    record = new_record(
        fichier, "PDF (lecture texte, à vérifier)",
        date=parse_date(dates[0]) if dates else None,
        fournisseur=next((l.strip() for l in text.splitlines() if l.strip()), "")[:60],
        tps=tps, tvh=tvh, tvq=tvq, total=total or None,
        remarques=["Lu sans IA (pas de clé ANTHROPIC_API_KEY) : vérifier les montants"],
    )
    return [apply_tax_included(record) if included else record]


# --------------------------------------------------------------------------
# One document
# --------------------------------------------------------------------------

def load_cache() -> dict:
    if CACHE_PATH.exists():
        try:
            cache = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
            if cache.get("version") == CACHE_VERSION:
                return cache
        except json.JSONDecodeError:
            pass
    return {"version": CACHE_VERSION, "documents": {}}


def save_cache(cache: dict) -> None:
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")


def extract_document(path: Path, cache: dict, api_key: str | None) -> tuple[list[dict], list[str]]:
    """Return (records, warnings) for one document."""
    fichier = relative(path)
    suffix = path.suffix.lower()
    digest = file_hash(path)
    cached = cache["documents"].get(digest)
    if cached:
        records = [dict(r, fichier=fichier) for r in cached["depenses"]]
        return records, list(cached.get("avertissements") or [])

    warnings: list[str] = []
    records: list[dict] = []
    use_cache = False
    try:
        if suffix in TABLE_EXT:
            sheets = read_table_rows(path)
            found = False
            for name, rows in sheets:
                parsed = table_records(rows, fichier, name if len(sheets) > 1 else "")
                if parsed is not None:
                    found = True
                    records += parsed
            if not found:
                if api_key:
                    block = {"type": "text", "text": f"Contenu du tableau « {path.name} » :\n"
                                                     + table_as_text(sheets)}
                    records = records_from_claude(call_claude(block, api_key), fichier, "tableau (lu par IA)")
                    use_cache = True
                else:
                    warnings.append(f"{fichier} : colonnes TPS et TVQ introuvables, document non compté")
        elif suffix in PDF_EXT:
            if api_key:
                block = {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                        "data": base64.b64encode(path.read_bytes()).decode()}}
                result = call_claude(block, api_key)
                records = records_from_claude(result, fichier, "PDF (lu par IA)")
                use_cache = True
                if not result.get("document_lisible", True):
                    warnings.append(f"{fichier} : document illisible ou sans dépense")
                if not records:
                    print(f"DEBUG {fichier} : réponse brute = "
                          f"{json.dumps(result, ensure_ascii=False)[:4000]}", file=sys.stderr)
            else:
                records = pdf_text_records(path, fichier)
                if not records:
                    warnings.append(f"{fichier} : PDF non lu (ajouter la clé ANTHROPIC_API_KEY), document non compté")
        elif suffix in IMAGE_EXT:
            if api_key:
                result = call_claude(image_block(path), api_key)
                records = records_from_claude(result, fichier, "photo (lue par IA)")
                use_cache = True
                if not result.get("document_lisible", True):
                    warnings.append(f"{fichier} : photo illisible ou sans dépense")
            else:
                warnings.append(f"{fichier} : photo non lue (ajouter la clé ANTHROPIC_API_KEY), photo non comptée")
        else:
            warnings.append(f"{fichier} : format {suffix or '(sans extension)'} non pris en charge "
                            f"(utiliser xlsx, csv, pdf, jpg ou png)")
    except Exception as exc:  # one bad document must not block the others
        warnings.append(f"{fichier} : erreur de lecture ({exc}), document non compté")
        return [], warnings

    if records == [] and not warnings:
        warnings.append(f"{fichier} : aucune dépense avec taxes trouvée")
    if use_cache:
        cache["documents"][digest] = {
            "fichier": fichier,
            "depenses": records,
            "avertissements": warnings,
            "extrait_le": datetime.now().isoformat(timespec="seconds"),
        }
    return records, warnings


# --------------------------------------------------------------------------
# A quarter
# --------------------------------------------------------------------------

def document_files(folder: Path, recursive: bool) -> list[Path]:
    if not folder.is_dir():
        return []
    paths = folder.rglob("*") if recursive else folder.iterdir()
    return sorted(p for p in paths
                  if p.is_file() and p.name.lower() not in IGNORED_NAMES and not p.name.startswith("~$"))


def loose_files() -> list[Path]:
    """Documents directly in depenses/ or in a sub-folder that is not a quarter folder."""
    files = document_files(EXPENSES_DIR, recursive=False)
    if EXPENSES_DIR.is_dir():
        for sub in sorted(EXPENSES_DIR.iterdir()):
            if sub.is_dir() and not QUARTER_DIR_RE.match(sub.name):
                files += document_files(sub, recursive=True)
    return files


def quarter_folder(label: str) -> Path | None:
    if not EXPENSES_DIR.is_dir():
        return None
    for sub in EXPENSES_DIR.iterdir():
        if sub.is_dir() and sub.name.upper() == label.upper():
            return sub
    return None


def remove_duplicates(records: list[dict]) -> tuple[list[dict], list[str]]:
    """Drop the same invoice found in two documents (e.g. photographed twice).

    Two expenses are the same invoice when they have the same supplier and
    invoice number, or the same supplier, date and total. The first one is
    kept; the others are left out of the totals and reported."""
    seen: dict[tuple, str] = {}
    kept, warnings = [], []
    for r in records:
        supplier = normalize(r["fournisseur"]).replace(" ", "")[:12]
        invoice = re.sub(r"[^0-9a-z]", "", (r["facture"] or "").lower())
        keys = []
        if invoice:
            keys.append(("facture", supplier, invoice))
        if r["date"] and r["total"]:
            keys.append(("montant", supplier, r["date"], r["total"]))
        first = next((seen[k] for k in keys if k in seen and seen[k] != r["fichier"]), None)
        if first:
            invoice_note = f" (facture {r['facture']})" if r["facture"] else ""
            warnings.append(f"Doublon exclu : {r['fournisseur']} du {r['date'] or '?'}"
                            f"{invoice_note}"
                            f" se trouve dans {first} et {r['fichier']} : comptée une seule fois")
            continue
        for k in keys:
            seen.setdefault(k, r["fichier"])
        kept.append(r)
    return kept, warnings


TELUS_RE = re.compile(r"\btelus\b")
TELUS_LOAN_RE = re.compile(r"\bpret\b|\bfinancement\b|\bappareil\b")
TELUS_MONTHS = 3


def apply_telus_rules(records: list[dict]) -> tuple[list[dict], list[str]]:
    """Telus conventions for the business cell phone.

    - Monthly service bill: one bill is deposited per quarter and the same
      amount is paid every month, so its HT/TPS/TVQ/total are multiplied by 3
      (TPS/TVQ calculated from the total first when the bill shows none).
    - "Prêt téléphone" (phone loan / device financing, detected by "prêt",
      "financement" or "appareil" in the file name or description): the tax
      was paid only once, so it is counted as printed, never multiplied.
    If several Telus monthly bills are in the same quarter, none is multiplied
    (each month is then already there) and a warning says so."""
    warnings: list[str] = []
    service, out = [], []
    for r in records:
        if TELUS_RE.search(normalize(r["fournisseur"])):
            r = dict(r)
            if TELUS_LOAN_RE.search(normalize(f"{r['fichier']} {r['description']}")):
                r["remarques"] = r["remarques"] + [
                    "Prêt téléphone Telus : taxe payée une seule fois, comptée telle quelle (pas de × 3)"]
            else:
                service.append(r)
        out.append(r)
    if len(service) == 1:
        r = apply_tax_included(service[0])
        for k in ("montant_ht", "tps", "tvq", "total"):
            if r[k] is not None:
                r[k] = r2(r[k] * TELUS_MONTHS)
        r["remarques"] = r["remarques"] + [
            f"Facture Telus mensuelle × {TELUS_MONTHS} (même montant payé chaque mois du trimestre)"]
        r["methode"] += f", × {TELUS_MONTHS} mois"
    elif len(service) > 1:
        warnings.append(f"{len(service)} factures Telus mensuelles dans le trimestre : "
                        f"aucune multipliée par {TELUS_MONTHS}, chacune comptée une fois")
    return out, warnings


def load_quarter_expenses(label: str, api_key: str | None = None) -> dict:
    """Extract every expense of quarter `label` (e.g. "2026-T4").

    Returns {"trimestre", "depenses", "avertissements", "documents", "tps", "tvq", "montant_ht"}.
    """
    api_key = api_key if api_key is not None else os.environ.get("ANTHROPIC_API_KEY") or None
    cache = load_cache()
    records: list[dict] = []
    warnings: list[str] = []
    documents: list[str] = []

    folder = quarter_folder(label)
    for path in document_files(folder, recursive=True) if folder else []:
        recs, warns = extract_document(path, cache, api_key)
        documents.append(relative(path))
        warnings += warns
        for r in recs:
            q = quarter_of_iso(r["date"])
            if q and q != label.upper():
                r["remarques"] = r["remarques"] + [f"Date {r['date']} hors du trimestre {label}"]
        records += recs

    for path in loose_files():
        recs, warns = extract_document(path, cache, api_key)
        mine = [r for r in recs if quarter_of_iso(r["date"]) == label.upper()]
        if mine:
            documents.append(relative(path))
            records += mine
        undated = [r for r in recs if not r["date"]]
        if undated:
            warnings.append(f"{relative(path)} : {len(undated)} dépense(s) sans date, non comptée(s) "
                            f"(ranger le document dans depenses/{label}/ pour la compter)")
        if mine or not recs:
            warnings += [w for w in warns if w not in warnings]

    save_cache(cache)
    records, duplicate_notes = remove_duplicates(records)
    warnings += duplicate_notes
    records, telus_notes = apply_telus_rules(records)
    warnings += telus_notes
    for r in records:
        r["categorie"] = categorize_fournisseur(r["fournisseur"])
        if r["devise"] != "CAD":
            warnings.append(f"{r['fichier']} : {r['fournisseur']} en {r['devise']}, montants non convertis en $ CA")
    records.sort(key=lambda r: (r["date"] or "9999", r["fichier"]))
    return {
        "trimestre": label.upper(),
        "depenses": records,
        "avertissements": warnings,
        "documents": documents,
        "montant_ht": r2(sum(r["montant_ht"] or 0 for r in records)),
        "tps": r2(sum(r["tps"] for r in records)),
        "tvq": r2(sum(r["tvq"] for r in records)),
    }


def quarters_for_paths(paths: list[str]) -> set[str]:
    """Quarters affected by changed paths under depenses/ (used by the workflow)."""
    quarters: set[str] = set()
    cache = load_cache()
    api_key = os.environ.get("ANTHROPIC_API_KEY") or None
    for raw in paths:
        parts = Path(raw.strip()).parts
        if not parts or parts[0] != "depenses":
            continue
        quarter_part = next((p for p in parts[1:-1] if QUARTER_DIR_RE.match(p)), None)
        if quarter_part:
            quarters.add(quarter_part.upper())
            continue
        path = ROOT / raw.strip()
        if path.is_file() and path.name.lower() not in IGNORED_NAMES:
            recs, _ = extract_document(path, cache, api_key)
            quarters.update(q for q in (quarter_of_iso(r["date"]) for r in recs) if q)
    save_cache(cache)
    return quarters


def main() -> None:
    if len(sys.argv) != 2 or not re.match(r"^\d{4}-[Tt][1-4]$", sys.argv[1]):
        print("Usage : python scripts/extract_expenses.py AAAA-TN   (ex. 2026-T4)", file=sys.stderr)
        sys.exit(1)
    result = load_quarter_expenses(sys.argv[1].upper())
    for r in result["depenses"]:
        print(f"{r['date'] or '?':10}  {r['fournisseur'][:30]:30}  {r['categorie']:35}"
              f"  TPS {r['tps']:>9.2f}  TVQ {r['tvq']:>9.2f}  [{r['methode']}] {r['fichier']}")
    print(f"\nTotal {result['trimestre']} : TPS payée {result['tps']:.2f} $, TVQ payée {result['tvq']:.2f} $ "
          f"({len(result['depenses'])} dépenses, {len(result['documents'])} documents)")
    for w in result["avertissements"]:
        print(f"ATTENTION : {w}")


if __name__ == "__main__":
    main()
