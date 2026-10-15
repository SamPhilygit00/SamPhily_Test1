#!/usr/bin/env python3
"""Build the quarterly TPS/TVQ declaration workbook from a Shopify orders JSON dump.

Reproduces the "Section A — Revenus" sheet:

  row 6   TAUX EN VIGUEUR
  row 7   Taux TPS (fédéral) 5,00 %   |   Taux TVQ (provincial) 9,975 %
  row 9   SECTION A — REVENUS (ventes et services facturés)
  row 10  Case | Description | Montant HT (Brut) ($) | TPS perçue ($) | TVQ perçue ($)
  row 11  101     Vente - En ligne (shopify): total sales before tax, and the
                  TPS/TVQ those sales represent at the rates of row 7
  row 12  105/205 Taxes perçues (clients): sales on which tax was charged, and
                  the TPS/TVQ actually collected

E12 and F12 are the "tps" and "tvq" of the "TOTAL PÉRIODE" row of the
"Suivi ventes" workbook built from the same orders (same per-day rounding),
so the two files always agree to the cent.

When the quarter's expenses are given (see extract_expenses.py), two more
sections follow:

  SECTION B — DÉPENSES   TPS/TVQ paid on purchases and fees (CTI / RTI),
                         detailed one expense per row in the "Dépenses" sheet
  SECTION C — SOLDE      TPS(Shopify) − TPS(dépenses) and the same for TVQ,
                         with a message: amount to pay to Revenu Québec when
                         the result is > 0, refund to claim when it is < 0

Usage:
  python scripts/build_tax_declaration_xlsx.py [--input data/orders/2026-T4.json] [--output path.xlsx]

When the input is named after a quarter (2026-T4.json), the expenses of
depenses/2026-T4/ are included (use --sans-depenses to leave them out).

Defaults: --input is the most recent extraction JSON in data/orders/,
--output is data/orders/<same-name>_declaration_tps_tvq.xlsx.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from build_sales_tracking_xlsx import (
    DATA_DIR,
    find_latest_input,
    load_orders,
    period_totals,
    split_taxes,
)

TPS_RATE = 0.05
TVQ_RATE = 0.09975

MONEY_FORMAT = '#,##0.00 "$"'
TPS_RATE_FORMAT = "0.00%"
TVQ_RATE_FORMAT = "0.000%"

TEAL_FILL = PatternFill("solid", fgColor="FF0E6E8C")
SECTION_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="FF2E75B6")
RATE_FILL = PatternFill("solid", fgColor="FFDCE6F1")
CASE_FILL = PatternFill("solid", fgColor="FFDCE6F1")
INPUT_FILL = PatternFill("solid", fgColor="FFFFFF00")
TAX_FILL = PatternFill("solid", fgColor="FFEEF4FB")

THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WHITE_TITLE = Font(name="Arial", bold=True, size=14, color="FFFFFFFF")
WHITE_HEADER = Font(name="Arial", bold=True, size=13, color="FFFFFFFF")
TEXT = Font(name="Arial", size=12)
BLUE_BOLD = Font(name="Arial", bold=True, size=13, color="FF0000FF")
BLUE = Font(name="Arial", size=12, color="FF0000FF")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


def declaration_amounts(orders: list[dict]) -> dict[str, float]:
    totals = period_totals(orders)
    # Sales on which TPS or TVQ was charged (orders shipped outside Québec /
    # Canada can have no tax), summed with the same per-day rounding.
    taxed = [o for o in orders if any(split_taxes(o))]
    taxed_sales = period_totals(taxed)["mt_brut"] if taxed else 0.0
    return {
        "ventes": totals["mt_brut"],
        "tps_theorique": round(totals["mt_brut"] * TPS_RATE, 2),
        "tvq_theorique": round(totals["mt_brut"] * TVQ_RATE, 2),
        "ventes_taxees": taxed_sales,
        "tps": totals["tps"],
        "tvq": totals["tvq"],
    }


def fr_money(value: float) -> str:
    """1234.5 -> '1 234,50 $' (French-Canadian format, for messages)."""
    return f"{value:,.2f}".replace(",", "\u00a0").replace(".", ",") + "\u00a0$"


def tax_balance(amounts: dict[str, float], expenses: dict | None) -> dict:
    """TPS/TVQ collected (Shopify) minus TPS/TVQ paid (expenses), with messages."""
    paid_tps = (expenses or {}).get("tps", 0.0)
    paid_tvq = (expenses or {}).get("tvq", 0.0)
    result = {"tps_percue": amounts["tps"], "tvq_percue": amounts["tvq"],
              "tps_payee": paid_tps, "tvq_payee": paid_tvq,
              "tps_nette": round(amounts["tps"] - paid_tps, 2),
              "tvq_nette": round(amounts["tvq"] - paid_tvq, 2)}
    messages = []
    for tax in ("tps", "tvq"):
        net = result[f"{tax}_nette"]
        name = tax.upper()
        if net > 0:
            messages.append(("payer", f"{name} : {fr_money(net)} à payer à Revenu Québec"))
        elif net < 0:
            messages.append(("reclamer", f"{name} : {fr_money(-net)} à réclamer à Revenu Québec (remboursement)"))
        else:
            messages.append(("zero", f"{name} : aucun montant à payer ni à réclamer"))
    result["messages"] = messages
    return result


def _cell(ws, ref: str, value=None, font=TEXT, fill=None, align=CENTER, fmt=None, border=True):
    cell = ws[ref]
    if value is not None:
        cell.value = value
    cell.font = font
    cell.alignment = align
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = BORDER
    return cell


def _banner(ws, row: int, text: str, fill: PatternFill) -> None:
    ws.merge_cells(f"C{row}:F{row}")
    _cell(ws, f"C{row}", text, font=WHITE_TITLE, fill=fill)
    for col in "DEF":
        ws[f"{col}{row}"].fill = fill
        ws[f"{col}{row}"].border = BORDER
    ws.row_dimensions[row].height = 24


def build_declaration(orders: list[dict], period: str = "", warnings: list[str] | None = None,
                      expenses: dict | None = None) -> Workbook:
    amounts = declaration_amounts(orders)

    wb = Workbook()
    ws = wb.active
    ws.title = "Déclaration TPS-TVQ"
    ws.sheet_view.showGridLines = True

    for col, width in {"A": 4, "B": 12, "C": 30, "D": 21, "E": 29, "F": 29}.items():
        ws.column_dimensions[col].width = width

    _cell(ws, "C2", "Déclaration TPS / TVQ", font=Font(name="Arial", bold=True, size=16),
          align=LEFT, border=False)
    if period:
        _cell(ws, "C3", f"Période : {period}", font=TEXT, align=LEFT, border=False)
    if warnings:
        _cell(ws, "C4", "ATTENTION — données incomplètes : " + " ".join(warnings),
              font=Font(name="Arial", bold=True, size=11, color="FFFF0000"), align=LEFT, border=False)

    _banner(ws, 6, "TAUX EN VIGUEUR", TEAL_FILL)
    _cell(ws, "C7", "Taux TPS (fédéral)", align=LEFT)
    _cell(ws, "D7", TPS_RATE, font=BLUE_BOLD, fill=RATE_FILL, fmt=TPS_RATE_FORMAT)
    _cell(ws, "E7", "Taux TVQ (provincial)", align=LEFT)
    _cell(ws, "F7", TVQ_RATE, font=BLUE_BOLD, fill=RATE_FILL, fmt=TVQ_RATE_FORMAT)

    _banner(ws, 9, "SECTION A — REVENUS (ventes et services facturés)", SECTION_FILL)

    headers = {"B": "Case", "C": "Description", "D": "Montant HT (Brut) ($)",
               "E": "TPS perçue ($)", "F": "TVQ perçue ($)"}
    for col, text in headers.items():
        _cell(ws, f"{col}10", text, font=WHITE_HEADER, fill=HEADER_FILL)
    ws.row_dimensions[10].height = 36

    rows = [
        (11, "101", "Vente -  En ligne (shopify)",
         amounts["ventes"], amounts["tps_theorique"], amounts["tvq_theorique"]),
        (12, "105/205", "Taxes perçues (clients)",
         amounts["ventes_taxees"], amounts["tps"], amounts["tvq"]),
    ]
    for row, case, label, montant, tps, tvq in rows:
        _cell(ws, f"B{row}", case, fill=CASE_FILL if row == 11 else None)
        _cell(ws, f"C{row}", label, align=LEFT)
        _cell(ws, f"D{row}", montant, font=BLUE, fill=INPUT_FILL, fmt=MONEY_FORMAT)
        _cell(ws, f"E{row}", tps, fill=TAX_FILL, fmt=MONEY_FORMAT)
        _cell(ws, f"F{row}", tvq, fill=TAX_FILL, fmt=MONEY_FORMAT)
        ws.row_dimensions[row].height = 18

    _cell(ws, "C14", "Ligne 101 : TPS/TVQ calculées sur le montant HT aux taux en vigueur (indicatif).",
          font=Font(name="Arial", italic=True, size=9, color="FF595959"), align=LEFT, border=False)
    _cell(ws, "C15", "Ligne 105/205 : TPS/TVQ réellement perçues = TOTAL PÉRIODE du fichier Suivi ventes.",
          font=Font(name="Arial", italic=True, size=9, color="FF595959"), align=LEFT, border=False)

    if expenses is not None:
        _expenses_sections(ws, amounts, expenses)
        _expenses_sheet(wb, expenses)

    for sheet in wb.worksheets:  # print each sheet on one page width
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1 if sheet is ws else 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    return wb


NOTE = Font(name="Arial", italic=True, size=9, color="FF595959")
RED_BOLD = Font(name="Arial", bold=True, size=13, color="FFC00000")
GREEN_BOLD = Font(name="Arial", bold=True, size=13, color="FF00703C")
BOLD = Font(name="Arial", bold=True, size=12)
PAY_FILL = PatternFill("solid", fgColor="FFFDE9E7")
CLAIM_FILL = PatternFill("solid", fgColor="FFE6F4EA")
TOTAL_FILL = PatternFill("solid", fgColor="FFFFF2CC")


def _expenses_sections(ws, amounts: dict[str, float], expenses: dict) -> None:
    label = expenses.get("trimestre", "")
    count = len(expenses.get("depenses", []))

    _banner(ws, 17, "SECTION B — DÉPENSES (taxes payées sur les achats et frais)", SECTION_FILL)
    headers = {"B": "Case", "C": "Description", "D": "Montant HT ($)",
               "E": "TPS payée ($)", "F": "TVQ payée ($)"}
    for col, text in headers.items():
        _cell(ws, f"{col}18", text, font=WHITE_HEADER, fill=HEADER_FILL)
    ws.row_dimensions[18].height = 36
    _cell(ws, "B19", "CTI / RTI", fill=CASE_FILL)
    _cell(ws, "C19", f"Dépenses ({count} pièce{'s' if count > 1 else ''})", align=LEFT)
    _cell(ws, "D19", expenses.get("montant_ht", 0.0), font=BLUE, fill=INPUT_FILL, fmt=MONEY_FORMAT)
    _cell(ws, "E19", expenses.get("tps", 0.0), fill=TAX_FILL, fmt=MONEY_FORMAT)
    _cell(ws, "F19", expenses.get("tvq", 0.0), fill=TAX_FILL, fmt=MONEY_FORMAT)
    note = ("Détail ligne par ligne dans l'onglet « Dépenses »." if count else
            f"Aucune dépense : déposer les factures et reçus dans le dossier depenses/{label}/ du dépôt.")
    _cell(ws, "C20", note, font=NOTE, align=LEFT, border=False)

    balance = tax_balance(amounts, expenses)
    _banner(ws, 22, "SECTION C — SOLDE À PAYER OU À RÉCLAMER", SECTION_FILL)
    for col, text in {"B": "", "C": "Description", "D": "", "E": "TPS ($)", "F": "TVQ ($)"}.items():
        _cell(ws, f"{col}23", text, font=WHITE_HEADER, fill=HEADER_FILL)
    ws.row_dimensions[23].height = 24
    rows = [
        (24, "Taxes perçues des clients (Shopify, ligne 105/205)", balance["tps_percue"], balance["tvq_percue"], TEXT),
        (25, "Moins : taxes payées sur les dépenses", -balance["tps_payee"], -balance["tvq_payee"], TEXT),
        (26, "Taxe nette (perçue − payée)", balance["tps_nette"], balance["tvq_nette"], BOLD),
    ]
    for row, text, tps, tvq, font in rows:
        fill = TOTAL_FILL if row == 26 else None
        _cell(ws, f"B{row}", "", fill=fill)
        ws.merge_cells(f"C{row}:D{row}")
        _cell(ws, f"C{row}", text, font=font, align=LEFT, fill=fill)
        _cell(ws, f"D{row}", fill=fill)
        _cell(ws, f"E{row}", tps, font=font, fill=fill or TAX_FILL, fmt=MONEY_FORMAT)
        _cell(ws, f"F{row}", tvq, font=font, fill=fill or TAX_FILL, fmt=MONEY_FORMAT)
        ws.row_dimensions[row].height = 18

    for row, (kind, message) in zip((28, 29), balance["messages"]):
        ws.merge_cells(f"B{row}:F{row}")
        font, fill = {"payer": (RED_BOLD, PAY_FILL), "reclamer": (GREEN_BOLD, CLAIM_FILL)}.get(kind, (BOLD, None))
        _cell(ws, f"B{row}", message, font=font, fill=fill, align=LEFT)
        for col in "CDEF":
            ws[f"{col}{row}"].border = BORDER
            if fill:
                ws[f"{col}{row}"].fill = fill
        ws.row_dimensions[row].height = 24

    row = 31
    for warning in expenses.get("avertissements", []):
        _cell(ws, f"C{row}", "ATTENTION — " + warning,
              font=Font(name="Arial", bold=True, size=10, color="FFFF0000"), align=LEFT, border=False)
        row += 1


def _expenses_sheet(wb: Workbook, expenses: dict) -> None:
    ws = wb.create_sheet("Dépenses")
    columns = [("Fichier", 34), ("Date", 12), ("Fournisseur", 28), ("N° facture", 16), ("Description", 30),
               ("Montant HT ($)", 15), ("TPS payée ($)", 14), ("TVQ payée ($)", 14), ("Total ($)", 14),
               ("Lu par", 22), ("Remarques", 50)]
    for index, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=title)
        cell.font = Font(name="Arial", bold=True, size=11, color="FFFFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"

    records = expenses.get("depenses", [])
    for row, r in enumerate(records, start=2):
        values = [r["fichier"], r["date"], r["fournisseur"], r["facture"], r["description"],
                  r["montant_ht"], r["tps"], r["tvq"], r["total"], r["methode"], " ; ".join(r["remarques"])]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = Font(name="Arial", size=10, color="FFC00000" if col == 11 and value else None)
            cell.border = BORDER
            if 6 <= col <= 9:
                cell.number_format = MONEY_FORMAT

    total_row = len(records) + 2
    ws.cell(row=total_row, column=1, value="TOTAL")
    for col, key in ((6, "montant_ht"), (7, "tps"), (8, "tvq")):
        ws.cell(row=total_row, column=col, value=expenses.get(key, 0.0)).number_format = MONEY_FORMAT
    for col in range(1, len(columns) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.font = Font(name="Arial", bold=True, size=11)
        cell.fill = TOTAL_FILL
        cell.border = BORDER


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=None)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--sans-depenses", action="store_true", help="ne pas inclure les dépenses")
    args = parser.parse_args()

    input_path = args.input or find_latest_input()
    orders = load_orders(input_path)
    output_path = args.output or DATA_DIR / f"{input_path.stem}_declaration_tps_tvq.xlsx"

    expenses = None
    if not args.sans_depenses and re.match(r"^\d{4}-T[1-4]$", input_path.stem):
        from extract_expenses import load_quarter_expenses
        expenses = load_quarter_expenses(input_path.stem)
        for _, message in tax_balance(declaration_amounts(orders), expenses)["messages"]:
            print(message)

    build_declaration(orders, period=input_path.stem, expenses=expenses).save(output_path)
    print(f"Wrote {output_path} from {len(orders)} orders in {input_path}")


if __name__ == "__main__":
    main()
